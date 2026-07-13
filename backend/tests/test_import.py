"""Iteration 5: Bulk Excel/CSV Import tests."""
import os
import io
import pytest
import requests
from datetime import datetime, timezone

BASE_URL = os.environ.get("REACT_APP_BACKEND_URL", "https://cashflow-hub-444.preview.emergentagent.com").rstrip("/")
TOKEN = os.environ.get("TEST_SESSION_TOKEN", "test_session_import_1783953951604")

RESOURCES = ["expenses", "incomes", "checks", "promissory-notes"]

EXPECTED_HEADERS = {
    "expenses": ["category", "description", "amount", "date"],
    "incomes": ["source", "description", "amount", "date"],
    "checks": ["type", "party", "amount", "due_date", "bank_name", "check_number", "status", "notes"],
    "promissory-notes": ["type", "party", "amount", "due_date", "status", "notes"],
}


@pytest.fixture(scope="module")
def client():
    s = requests.Session()
    s.headers.update({"Authorization": f"Bearer {TOKEN}"})
    return s


@pytest.fixture(scope="module", autouse=True)
def cleanup(client):
    yield
    # Cleanup any TEST_ prefixed records
    for res in ["expenses", "incomes", "checks", "promissory-notes"]:
        try:
            r = client.get(f"{BASE_URL}/api/{res}")
            if r.status_code == 200:
                for item in r.json():
                    desc = (item.get("description") or item.get("party") or item.get("source") or "")
                    if "TEST_IMPORT" in str(desc) or "TEST_IMPORT" in str(item.get("category", "")):
                        client.delete(f"{BASE_URL}/api/{res}/{item['id']}")
        except Exception:
            pass


# --- Template download ---
@pytest.mark.parametrize("resource", RESOURCES)
def test_template_download(client, resource):
    r = client.get(f"{BASE_URL}/api/import/{resource}/template")
    assert r.status_code == 200, r.text
    ct = r.headers.get("content-type", "")
    assert "spreadsheetml" in ct or "xlsx" in ct, f"content-type: {ct}"
    assert len(r.content) > 100

    # Verify xlsx contains headers
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    assert len(rows) >= 2, "Template must have header + sample row"
    header_row = [str(c).strip().lower() for c in rows[0]]
    for h in EXPECTED_HEADERS[resource]:
        assert h in header_row, f"Missing header {h} for {resource}: got {header_row}"


def test_template_unknown_resource(client):
    r = client.get(f"{BASE_URL}/api/import/unknown/template")
    assert r.status_code == 404


# --- CSV import: basic with Turkish, BOM, comma delim ---
def test_import_expenses_csv_utf8_bom(client):
    csv_text = "category,description,amount,date\nKira,TEST_IMPORT Şubat ç ğ ş,12500.50,2026-02-01\nMalzeme,TEST_IMPORT Mart,3200,15/03/2026\n"
    content = ("\ufeff" + csv_text).encode("utf-8")
    files = {"file": ("test.csv", content, "text/csv")}
    r = client.post(f"{BASE_URL}/api/import/expenses", files=files)
    assert r.status_code == 200, r.text
    data = r.json()
    assert data["inserted"] == 2, data
    assert data["total_rows"] == 2
    assert data["errors"] == []

    # Verify persistence via GET
    listing = client.get(f"{BASE_URL}/api/expenses").json()
    descs = [e.get("description") for e in listing]
    assert any("Şubat" in (d or "") for d in descs), "Turkish char lost"


def test_import_incomes_csv_semicolon(client):
    csv_text = "source;description;amount;date\nSatış;TEST_IMPORT gelir;45000,00;15.01.2026\n"
    files = {"file": ("test.csv", csv_text.encode("utf-8"), "text/csv")}
    r = client.post(f"{BASE_URL}/api/import/incomes", files=files)
    assert r.status_code == 200, r.text
    data = r.json()
    assert data["inserted"] == 1, data


# --- Missing columns → 400 ---
def test_import_missing_required_columns(client):
    csv_text = "description,amount\nfoo,100\n"
    files = {"file": ("bad.csv", csv_text.encode("utf-8"), "text/csv")}
    r = client.post(f"{BASE_URL}/api/import/expenses", files=files)
    assert r.status_code == 400
    detail = r.json().get("detail", "")
    assert "category" in detail.lower() or "eksik" in detail.lower()


# --- Invalid amount / date → errors[] but valid inserted ---
def test_import_partial_errors(client):
    csv_text = (
        "category,description,amount,date\n"
        "Kira,TEST_IMPORT ok,1000,2026-02-01\n"
        "Kira,TEST_IMPORT bad amount,abc,2026-02-01\n"
        "Kira,TEST_IMPORT bad date,500,not-a-date\n"
    )
    files = {"file": ("mix.csv", csv_text.encode("utf-8"), "text/csv")}
    r = client.post(f"{BASE_URL}/api/import/expenses", files=files)
    assert r.status_code == 200, r.text
    data = r.json()
    assert data["inserted"] == 1
    assert len(data["errors"]) == 2
    rows = {e["row"] for e in data["errors"]}
    assert rows == {3, 4}


# --- XLSX import with datetime cells ---
def test_import_expenses_xlsx(client):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.append(["category", "description", "amount", "date"])
    ws.append(["Kira", "TEST_IMPORT xlsx ğüş", 9999.99, datetime(2026, 4, 10)])
    ws.append(["Diğer", "TEST_IMPORT xlsx 2", 100, "20.05.2026"])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    files = {"file": ("data.xlsx", buf.read(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    r = client.post(f"{BASE_URL}/api/import/expenses", files=files)
    assert r.status_code == 200, r.text
    data = r.json()
    assert data["inserted"] == 2, data

    # Verify date parsed to ISO
    listing = client.get(f"{BASE_URL}/api/expenses").json()
    found = [e for e in listing if e.get("description") == "TEST_IMPORT xlsx ğüş"]
    assert found and found[0]["date"] == "2026-04-10", found


# --- Enum validation for checks ---
def test_import_checks_enum(client):
    csv_text = (
        "type,party,amount,due_date,bank_name,check_number,status,notes\n"
        "issued,TEST_IMPORT chk1,15000,2026-03-15,Garanti,12345,pending,\n"
        "bogus,TEST_IMPORT chk_bad,15000,2026-03-15,Garanti,12345,pending,\n"
        "received,TEST_IMPORT chk_badstatus,15000,2026-03-15,Garanti,12345,invalid,\n"
    )
    files = {"file": ("chk.csv", csv_text.encode("utf-8"), "text/csv")}
    r = client.post(f"{BASE_URL}/api/import/checks", files=files)
    assert r.status_code == 200, r.text
    data = r.json()
    assert data["inserted"] == 1, data
    assert len(data["errors"]) == 2


def test_import_notes_enum(client):
    csv_text = (
        "type,party,amount,due_date,status,notes\n"
        "received,TEST_IMPORT note1,8500,2026-03-20,pending,\n"
        "issued,TEST_IMPORT note2,500,2026-03-21,overdue,\n"
        "received,TEST_IMPORT note_bad,500,2026-03-21,wrongstatus,\n"
    )
    files = {"file": ("n.csv", csv_text.encode("utf-8"), "text/csv")}
    r = client.post(f"{BASE_URL}/api/import/promissory-notes", files=files)
    assert r.status_code == 200, r.text
    data = r.json()
    assert data["inserted"] == 2, data
    assert len(data["errors"]) == 1


# --- File size limit ---
def test_import_file_too_large(client):
    # 6MB junk file
    big = b"x" * (6 * 1024 * 1024)
    files = {"file": ("big.csv", big, "text/csv")}
    r = client.post(f"{BASE_URL}/api/import/expenses", files=files)
    assert r.status_code == 413


# --- Auth required ---
def test_import_requires_auth():
    s = requests.Session()
    r = s.get(f"{BASE_URL}/api/import/expenses/template")
    assert r.status_code == 401
    r = s.post(f"{BASE_URL}/api/import/expenses", files={"file": ("a.csv", b"a,b\n1,2\n", "text/csv")})
    assert r.status_code == 401


# --- Data isolation: imported records have user_id = data_owner_id ---
def test_import_data_isolation(client):
    csv_text = "category,description,amount,date\nIsolate,TEST_IMPORT iso,50,2026-02-01\n"
    files = {"file": ("iso.csv", csv_text.encode("utf-8"), "text/csv")}
    r = client.post(f"{BASE_URL}/api/import/expenses", files=files)
    assert r.status_code == 200 and r.json()["inserted"] == 1
    # As logged in user, we should see the record
    listing = client.get(f"{BASE_URL}/api/expenses").json()
    assert any(e.get("description") == "TEST_IMPORT iso" for e in listing)
    # Without auth we shouldn't
    r2 = requests.get(f"{BASE_URL}/api/expenses")
    assert r2.status_code == 401
